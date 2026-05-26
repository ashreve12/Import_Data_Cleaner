"""Merge cleaned import workbooks into a destination master workbook."""

from __future__ import annotations

import re
from copy import copy
from datetime import datetime
from pathlib import Path
from typing import Sequence

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries


def _find_source_ws(wb, sheet_name: str | None):
    if sheet_name:
        if sheet_name in wb.sheetnames:
            return wb[sheet_name]
        raise ValueError(f"Source sheet '{sheet_name}' not found. Available: {wb.sheetnames}")
    return wb.active


def _find_dest_ws(wb, sheet_name: str | None):
    if sheet_name:
        if sheet_name in wb.sheetnames:
            return wb[sheet_name]
        raise ValueError(f"Destination sheet '{sheet_name}' not found. Available: {wb.sheetnames}")
    return wb.active


def _get_table_bounds(ws, table_name: str | None) -> tuple[int, int, int, int] | None:
    if not ws.tables:
        return None
    table = ws.tables.get(table_name) if table_name else next(iter(ws.tables.values()), None)
    if not table:
        return None
    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    return min_col, min_row, max_col, max_row


def _find_header_row(ws, table_name: str | None = None) -> tuple[int, int]:
    """Return (header_row, last_data_row). Uses table bounds if present, else row 1."""
    bounds = _get_table_bounds(ws, table_name)
    if bounds:
        _, header_row, _, last_data_row = bounds
        return header_row, last_data_row
    return 1, ws.max_row


def _clean_value(val):
    """Strip the time component from midnight datetimes so Excel shows only the date."""
    if isinstance(val, datetime) and val.hour == 0 and val.minute == 0 and val.second == 0:
        return val.date()
    return val


def _normalize(val) -> str:
    return str(val).strip().lower() if val is not None else ""


def _headers_from_tuple(row_tuple) -> dict[str, int]:
    """Build normalized header → 1-based column index from an iter_rows tuple."""
    headers: dict[str, int] = {}
    for col_idx, val in enumerate(row_tuple, start=1):
        key = _normalize(val)
        if key and key not in headers:
            headers[key] = col_idx
    return headers


def _headers_from_ws(ws, row: int) -> dict[str, int]:
    """Build normalized header → 1-based column index by cell access."""
    headers: dict[str, int] = {}
    for col_idx in range(1, ws.max_column + 1):
        key = _normalize(ws.cell(row=row, column=col_idx).value)
        if key and key not in headers:
            headers[key] = col_idx
    return headers


def _adjust_formula_row(formula: str, delta: int) -> str:
    def _replace(m: re.Match) -> str:
        col_part, row_dollar, row_num = m.group(1), m.group(2), int(m.group(3))
        if row_dollar:
            return f"{col_part}{row_dollar}{row_num}"
        return f"{col_part}{row_num + delta}"
    return re.sub(r"(\$?[A-Za-z]+)(\$?)(\d+)", _replace, formula)


def _append_source(
    source_path: Path,
    source_sheet_name: str | None,
    dest_ws,
    dest_headers: dict[str, int],
    formula_cols: dict[int, str],
    last_data_row: int,
    header_scan_limit: int = 30,
) -> tuple[int, int]:
    """Stream source rows (read-only) into dest_ws. Returns (rows_added, matched_columns)."""
    source_wb = load_workbook(source_path, read_only=True, data_only=True)
    try:
        source_ws = _find_source_ws(source_wb, source_sheet_name)

        # Buffer first N rows to find the header.
        # Rule: use the first row that has >= min_header_cols non-empty cells.
        # This skips sparse title rows (e.g. Module Final has a 1-cell title in
        # row 1; the real headers are in row 2) without being tricked by data
        # rows that are wider than the header row.
        min_header_cols = 3
        prefix: list[tuple] = []
        header_idx: int | None = None
        for row in source_ws.iter_rows(min_row=1, max_row=header_scan_limit, values_only=True):
            prefix.append(row)
            if header_idx is None and sum(1 for v in row if v is not None) >= min_header_cols:
                header_idx = len(prefix) - 1

        if header_idx is None:
            header_idx = 0  # nothing dense enough found — fall back to row 1

        source_headers = _headers_from_tuple(prefix[header_idx] if prefix else ())

        col_map: dict[int, int] = {}  # dest_col → source_col (1-based)
        for header, dest_col in dest_headers.items():
            if header in source_headers:
                col_map[dest_col] = source_headers[header]

        if not col_map:
            raise ValueError(
                f"No matching column headers found between source '{source_ws.title}' "
                f"({len(source_headers)} cols) and destination '{dest_ws.title}' "
                f"({len(dest_headers)} cols). "
                f"Source: {sorted(source_headers)}. "
                f"Destination: {sorted(dest_headers)}."
            )

        dest_max_col = dest_ws.max_column
        rows_added = 0

        # Data rows = prefix rows after the header + the rest of the file streamed
        def _data_rows():
            for i, row in enumerate(prefix):
                if i > header_idx:
                    yield row
            yield from source_ws.iter_rows(min_row=len(prefix) + 1, values_only=True)

        for row_values in _data_rows():
            if all(v is None for v in row_values):
                continue

            dest_row = last_data_row + rows_added + 1
            delta = rows_added + 1

            for dest_col in range(1, dest_max_col + 1):
                template_cell = dest_ws.cell(row=last_data_row, column=dest_col)
                new_cell = dest_ws.cell(row=dest_row, column=dest_col)

                if dest_col in col_map:
                    src_col = col_map[dest_col]
                    raw = row_values[src_col - 1] if src_col <= len(row_values) else None
                    new_cell.value = _clean_value(raw)
                elif dest_col in formula_cols:
                    new_cell.value = _adjust_formula_row(formula_cols[dest_col], delta)

                if template_cell.has_style:
                    new_cell.font = copy(template_cell.font)
                    new_cell.fill = copy(template_cell.fill)
                    new_cell.border = copy(template_cell.border)
                    new_cell.alignment = copy(template_cell.alignment)
                    new_cell.number_format = template_cell.number_format

            template_height = dest_ws.row_dimensions[last_data_row].height
            if template_height:
                dest_ws.row_dimensions[dest_row].height = template_height

            rows_added += 1

        return rows_added, len(col_map)
    finally:
        source_wb.close()


def merge_all_workbooks(
    source_paths: Sequence[Path],
    dest_path: Path,
    output_path: Path,
    dest_sheet_name: str | None = None,
    source_sheet_name: str | None = None,
    dest_table_name: str | None = None,
) -> dict:
    """Merge multiple source workbooks into dest, loading dest only once."""
    dest_wb = load_workbook(dest_path)
    try:
        dest_ws = _find_dest_ws(dest_wb, dest_sheet_name)
        dest_header_row, last_data_row = _find_header_row(dest_ws, dest_table_name)
        found_table_name = (
            (dest_table_name or next(iter(dest_ws.tables), None))
            if dest_ws.tables else None
        )
        dest_headers = _headers_from_ws(dest_ws, row=dest_header_row)

        formula_cols: dict[int, str] = {}
        for dest_col in range(1, dest_ws.max_column + 1):
            cell = dest_ws.cell(row=last_data_row, column=dest_col)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formula_cols[dest_col] = cell.value

        total_rows = 0
        total_matched = 0
        current_last_row = last_data_row

        for src_path in source_paths:
            rows_added, matched = _append_source(
                src_path, source_sheet_name, dest_ws,
                dest_headers, formula_cols, current_last_row,
            )
            total_rows += rows_added
            total_matched = matched  # same dest headers each time
            current_last_row += rows_added

        if found_table_name and total_rows > 0:
            table = dest_ws.tables[found_table_name]
            min_col, min_row, max_col, _ = range_boundaries(table.ref)
            table.ref = (
                f"{get_column_letter(min_col)}{min_row}"
                f":{get_column_letter(max_col)}{current_last_row}"
            )

        dest_wb.save(output_path)
    finally:
        dest_wb.close()

    return {
        "destination_sheet": dest_ws.title,
        "destination_table": found_table_name or "none",
        "matched_columns": total_matched,
        "formula_columns_adjusted": len(formula_cols),
        "rows_added": total_rows,
        "output_file": str(output_path),
    }


def merge_workbooks(
    source_path: Path,
    dest_path: Path,
    output_path: Path,
    dest_sheet_name: str | None = None,
    source_sheet_name: str | None = None,
    dest_table_name: str | None = None,
) -> dict:
    return merge_all_workbooks(
        source_paths=[source_path],
        dest_path=dest_path,
        output_path=output_path,
        dest_sheet_name=dest_sheet_name,
        source_sheet_name=source_sheet_name,
        dest_table_name=dest_table_name,
    )
